from __future__ import annotations

from collections import OrderedDict
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .database import TrainerDatabase


HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_FONT = Font(color="FFFFFFFF", bold=True)
BODY_FILL = PatternFill("solid", fgColor="FFFFFFFF")
BODY_FONT = Font(color="FF111827")
ROW_BORDER = Border(bottom=Side(style="thin", color="FFE5E7EB"))
DAY_BORDER = Border(top=Side(style="medium", color="FF94A3B8"))


def _local_datetime(value: str, timezone: ZoneInfo) -> datetime:
    return datetime.fromisoformat(value).astimezone(timezone).replace(tzinfo=None)


def _format_distance(value: float) -> str:
    return f"{value:g}".replace(".", ",") + " км"


def build_workout_export(
    database: TrainerDatabase, telegram_user_id: int, timezone: ZoneInfo
) -> bytes:
    grouped_strength: OrderedDict[int, dict] = OrderedDict()
    for row in database.get_export_sets(telegram_user_id):
        workout = grouped_strength.setdefault(
            int(row["workout_id"]),
            {
                "completed_at": row["completed_at"],
                "exercise_name": row["exercise_name"],
                "weights": [],
            },
        )
        workout["weights"].append(float(row["weight"]))

    activities = [
        {
            "completed_at": workout["completed_at"],
            "workout_id": workout_id,
            "exercise_name": workout["exercise_name"],
            "sets": workout["weights"],
        }
        for workout_id, workout in grouped_strength.items()
    ]
    activities.extend(
        {
            "completed_at": row["completed_at"],
            "workout_id": int(row["workout_id"]),
            "exercise_name": row["exercise_name"],
            "sets": [_format_distance(float(row["distance_km"]))],
        }
        for row in database.get_export_distances(telegram_user_id)
    )
    activities.sort(key=lambda item: (item["completed_at"], item["workout_id"]))

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "История"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(
        ["Дата", "Упражнение", "1-й подход", "2-й подход", "3-й подход"]
    )

    for activity in activities:
        completed_at = _local_datetime(activity["completed_at"], timezone)
        sets = list(activity["sets"][:3])
        sets.extend([None] * (3 - len(sets)))
        sheet.append(
            [completed_at.strftime("%d.%m.%Y"), activity["exercise_name"], *sets]
        )

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 26

    previous_date = None
    for row in sheet.iter_rows(min_row=2, max_col=5):
        current_date = row[0].value
        for cell in row:
            cell.fill = BODY_FILL
            cell.font = BODY_FONT
            cell.border = DAY_BORDER if current_date != previous_date else ROW_BORDER
            cell.alignment = Alignment(
                horizontal="left" if cell.column == 2 else "center",
                vertical="center",
            )
        row[0].number_format = "@"
        previous_date = current_date

    sheet.column_dimensions["A"].width = 15
    sheet.column_dimensions["B"].width = 30
    for column in ("C", "D", "E"):
        sheet.column_dimensions[column].width = 16

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
